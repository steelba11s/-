from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
from hashlib import sha256
from io import BytesIO
import json, re
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[2]
QA = Path(__file__).resolve().parent
OUT = ROOT / 'Материалы для защиты'
REF = Path('C:/Users/Artur/.codex/plugins/cache/openai-curated-remote/openai-templates/0.1.1/skills/artifact-template-system-design/assets/reference.docx')
OUT.mkdir(exist_ok=True)
ref = Document(REF)
NS = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
baseline_hash = sha256(REF.read_bytes()).hexdigest()

def contract():
    with ZipFile(REF) as z:
        inv={n:{'bytes':len(z.read(n)),'sha256':sha256(z.read(n)).hexdigest()} for n in z.namelist()}
    (QA/'reference-parts.json').write_text(json.dumps(inv,indent=2),encoding='utf-8')
    sect=ref.sections[0]
    text=f'''# Contract for the typing trainer defense guide
Reference: {REF}
SHA256: {baseline_hash}
Reference pages: 7; sections: 1. Reference rendered by Word COM after packaged renderer failed because bundled LibreOffice is unavailable on Windows. All seven page PNGs inspected.
Page dimensions EMU: {sect.page_width} x {sect.page_height}; margins top {sect.top_margin}, bottom {sect.bottom_margin}, left {sect.left_margin}, right {sect.right_margin}. Preserve sectPr including first-page behavior and header/footer distances.
Reference paragraph 22 is body: normal style, 1.25 line spacing, 5.5pt after, justified, Helvetica Neue, dark body color 233447. Reuse pPr and rPr.
Title: copy reference paragraph 9 with Title style and 22pt inherited size; black text for guide heading. Heading 1 copies reference paragraph 21, 13.5pt source style, keep-next. Minor teaching headings clone Heading 3 source pattern with black font. Code: new teaching component, monospaced 9pt, pale source-family fill, 1.1 line spacing.
Tables: clone source table style and explicit width geometry; dark blue 0B2D4B header and pale blue alternating rows, visible D9D9D9 borders; automatically expanding rows, 90 twip cell padding, repeat header. Figure slots replaced with real test-generated app images.
Editable slots: entire word/document.xml body (excluding sectPr) expanded into a beginner manual per user request. Replace proposal-only placeholders, dummy diagrams, metadata, sample links and example footnote anchors. Retain source package parts and styles; obsolete unreferenced assets stay opaque. Footer placeholder becomes guide name and PAGE field. Use source title, body, heading and table component family throughout expanded content.
Content order: title and learning plan; C# foundations; Avalonia foundations; project walkthrough; tests and run; defense speech and Q&A; source map. No invented author claims, backend APIs, server deployment or future features as existing.
Preserve-only: original styles, themes, numbering, section geometry, customXml, headers, footnotes and all original relationships and media; body relationship set may gain figures. All baseline package parts except body and footer text preserved byte-for-byte. Package inventory: reference-parts.json.
Final gate: no placeholders in visible body/footer; reference hash unchanged; page geometry and styles preserved; all final pages rendered and visually inspected; no clipped code, tables or orphan headings.
'''
    (QA/'artifact.md').write_text(text,encoding='utf-8')

contract()

def apply_proto(p, index):
    src=ref.paragraphs[index]
    if src._p.pPr is not None:
        if p._p.pPr is not None: p._p.remove(p._p.pPr)
        p._p.insert(0,deepcopy(src._p.pPr))
    for run in p.runs:
        if src.runs and src.runs[0]._r.rPr is not None:
            if run._r.rPr is not None: run._r.remove(run._r.rPr)
            run._r.insert(0,deepcopy(src.runs[0]._r.rPr))
    return p

def para(d,text='',kind='body'):
    idx={'body':22,'h1':21,'h2':35,'title':9}[kind]
    p=d.add_paragraph(text)
    apply_proto(p,idx)
    # Inherited numbering belongs to template headings, but guide numbers are explicit.
    if p._p.pPr is not None:
        n=p._p.pPr.find(qn('w:numPr'))
        if n is not None:p._p.pPr.remove(n)
    if kind!='body':
        p.paragraph_format.keep_with_next=True
        for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
    if kind=='h2':
        p.paragraph_format.space_before=Pt(9)
        p.paragraph_format.space_after=Pt(4)
    if kind=='h1':p.paragraph_format.space_before=Pt(14)
    p.paragraph_format.widow_control=True
    return p

def code(d,lines):
    for i,line in enumerate(lines):
        p=d.add_paragraph()
        p.paragraph_format.space_before=Pt(0)
        p.paragraph_format.space_after=Pt(0 if i<len(lines)-1 else 7)
        p.paragraph_format.line_spacing=1.1
        p.paragraph_format.left_indent=Pt(7)
        p.paragraph_format.right_indent=Pt(5)
        p.paragraph_format.keep_with_next=i<len(lines)-1
        p.paragraph_format.keep_together=True
        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F1F5F8');p._p.get_or_add_pPr().append(shade)
        r=p.add_run(line if line else ' ');r.font.name='Consolas';r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('233447')

def table(d,rows):
    rows=[row for row in rows if not all(re.fullmatch(r'[:\- ]+',c) for c in row)]
    t=d.add_table(rows=0,cols=len(rows[0]))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    total=7.1
    widths=[2.7,total-2.7] if len(rows[0])==2 else [total/len(rows[0])]*len(rows[0])
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for idx,row in enumerate(rows):
        cells=t.add_row().cells
        for j,(cell,value) in enumerate(zip(cells,row)):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.add_run(value);apply_proto(p,22)
            p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.1
            for r in p.runs:r.font.size=Pt(9);r.bold=idx==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if idx==0 else '233447')
            tcpr=cell._tc.get_or_add_tcPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'0B2D4B' if idx==0 else ('E6F0F8' if idx%2 else 'FFFFFF'));tcpr.append(shd)
            mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');mar.append(e)
            tcpr.append(mar)
        trpr=t.rows[-1]._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');trpr.append(ns)
        if idx==0:trpr.append(OxmlElement('w:tblHeader'))
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    t._tbl.tblPr.append(borders)
    d.add_paragraph().paragraph_format.space_after=Pt(2)

def figure(d,key):
    path=ROOT/'TypingTrainer.Tests/bin/Debug/net9.0/screenshots/trainer-1240x800.png'
    p=d.add_paragraph();p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(path),width=Inches(7.0))

def markdown(d,content,offset=0):
    lines=content.replace('».»','».').splitlines();i=0;n=offset
    while i<len(lines):
        line=lines[i]
        if not line.strip():i+=1;continue
        if line.startswith('```'):
            block=[];i+=1
            while i<len(lines) and not lines[i].startswith('```'):block.append(lines[i]);i+=1
            code(d,block)
        elif line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                rows.append([x.strip() for x in lines[i].strip('|').split('|')]);i+=1
            table(d,rows);continue
        elif line.startswith('# '):
            n+=1;para(d,f'{n} {line[2:]}','h1')
        elif line.startswith('## '):para(d,line[3:],'h2')
        elif line.startswith('@image '):figure(d,line.split()[1])
        else:para(d,line)
        i+=1
    return n

def build(key,title,filename):
    d=Document(REF)
    for child in list(d._element.body):
        if child.tag!=qn('w:sectPr'):d._element.body.remove(child)
    p=para(d,title,'title');p.paragraph_format.space_before=Pt(90)
    para(d,'Пособие для понимания кода и защиты проекта','title')
    p=para(d,'C# и Avalonia с нуля');p.paragraph_format.space_before=Pt(20)
    para(d,'Здесь разобрано, что делает приложение, где находятся его основные части и как объяснить их преподавателю. Сначала изучи базовые понятия, затем проследи действия пользователя по коду. Для защиты важнее понимать причины каждого шага, чем помнить все строки наизусть.')
    para(d,'Как пользоваться пособием','h2')
    para(d,'Первый проход — основы C# и Avalonia. Второй — устройство именно этого приложения с открытыми рядом исходниками. Третий — вопросы, речь и демонстрация. Короткие примеры кода можно читать как перевод обычного действия на язык программы.')
    para(d,'Если до защиты остался час, прочитай основной сценарий, архитектуру, ключевые алгоритмы и ограничения, затем проговори речь и выполни план демонстрации. После этого вернись к незнакомым словам в начальных разделах.')
    table(d,[['Сведения','Значение'],['Приложение',title],['Исходники', 'TypingTrainerApp.slnx'],['Дата проверки','11 сентября 2026'],['Автоматические тесты','15 пройдено']])
    d.add_page_break()
    common=(QA/'common.md').read_text(encoding='utf-8')
    content=(QA/(key+'.md')).read_text(encoding='utf-8')
    n=markdown(d,common)
    n=markdown(d,content,n)
    para(d,f'{n+1} Где проверить объяснения','h1')
    para(d,'Главный источник описания приложения — его исходники. Имена классов и методов в тексте позволяют найти соответствующий участок поиском в редакторе кода. Версии библиотек указаны по .csproj, результаты тестов — по запуску 11 сентября 2026 года. Учебные фрагменты отдельно названы учебными; основные алгоритмы приведены из проекта.')
    for x in ['TypingTrainer.Core/TypingSession.cs — состояние попытки, сравнение и формулы.','TypingTrainer.Core/AppDataStore.cs и Models.cs — сохранение и модели данных.','TypingTrainer/MainWindow.axaml и MainWindow.axaml.cs — окно и обработчики.','TypingTrainer.Tests/CoreTests.cs и WindowTests.cs — проверяемые сценарии.','typing-trainer-notebook.ipynb — анализ истории.']:para(d,x)
    para(d,'Официальные материалы для уточнения терминов','h2')
    sources=[('Microsoft Learn — обзор языка C#','https://learn.microsoft.com/en-us/dotnet/csharp/tour-of-csharp/'),('Microsoft Learn — асинхронное программирование','https://learn.microsoft.com/en-us/dotnet/csharp/asynchronous-programming/'),('Avalonia — контекст данных','https://docs.avaloniaui.net/docs/data-binding/data-context'),('Avalonia — синтаксис привязки','https://docs.avaloniaui.net/docs/data-binding/data-binding-syntax'),('Avalonia — MVVM','https://docs.avaloniaui.net/docs/fundamentals/the-mvvm-pattern')]
    for label,url in sources:
        p=para(d,label+'\n'+url);p.paragraph_format.space_after=Pt(5);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    # Replace visible footer metadata while retaining source section furniture.
    for section in d.sections:
        for footer in [section.footer]:
            for p in footer.paragraphs:
                if '[Organization Name]' in p.text:
                    p.text=title+'  |  '
                    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
                    for r in p.runs:r.font.name='Helvetica Neue';r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('5A7085')
    bio=BytesIO();d.save(bio)
    # Preserve opaque template parts byte-for-byte; only authoring slots may change.
    path=OUT/filename
    with ZipFile(REF) as original, ZipFile(BytesIO(bio.getvalue())) as generated, ZipFile(path,'w',ZIP_DEFLATED) as result:
        editable={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
        editable.update(n for n in original.namelist() if re.fullmatch(r'word/footer\d+\.xml',n))
        for info in original.infolist():result.writestr(info,generated.read(info.filename) if info.filename in editable and info.filename in generated.namelist() else original.read(info.filename))
        for name in generated.namelist():
            if name not in original.namelist():result.writestr(name,generated.read(name))
    with ZipFile(REF) as a,ZipFile(path) as b:
        changed=[name for name in a.namelist() if a.read(name)!=b.read(name)]
        assert set(changed)<=editable,changed
        roota=etree.fromstring(a.read('word/document.xml'));rootb=etree.fromstring(b.read('word/document.xml'))
        def signature(el):return (el.tag,dict(el.attrib),el.text,tuple(signature(x) for x in el))
        assert signature(roota.find('.//w:sectPr',NS))==signature(rootb.find('.//w:sectPr',NS))
    assert sha256(REF.read_bytes()).hexdigest()==baseline_hash
    print(json.dumps({'path':str(path),'paragraphs':len(d.paragraphs),'words':len(' '.join(p.text for p in d.paragraphs).split()),'changed_parts':changed},ensure_ascii=False))

build('trainer','Тренажёр набора текста','Тренажёр набора текста — подготовка к защите.docx')
